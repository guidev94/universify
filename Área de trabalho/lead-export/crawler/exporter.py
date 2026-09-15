import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


ALL_COLUMNS = {
    "name": "Nome",
    "email": "Email",
    "phone": "Telefone",
    "whatsapp": "WhatsApp",
    "wa_link": "WA_Link",
    "website": "Website",
    "score": "Score",
    "tier": "Tier",
    "status": "Status",
    "deal_value": "Valor Negócio",
    "deal_notes": "Notas",
    "source": "Origem",
    "category": "Categoria",
    "site_verified": "Site Verificado",
    "site_status": "Status do Site",
    "ssl_expired": "SSL Expirado",
    "copyright_year": "Ano Copyright",
    "is_responsive": "Responsivo",
    "cms": "CMS",
    "cnpj": "CNPJ",
    "partners": "Sócios",
    "has_ads": "Anuncia",
    "has_pixel": "Pixel",
    "has_analytics": "Analytics",
    "tech_score": "Tech Score",
    "company_value": "Capital Social",
    "legal_nature": "Nat. Juridica",
    "location": "Localização",
    "instagram": "Instagram",
    "facebook": "Facebook",
}

def export_csv(leads, filepath, selected_columns=None):
    if not selected_columns:
        selected_columns = list(ALL_COLUMNS.keys())

    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        
        # Header
        writer.writerow([ALL_COLUMNS[col] for col in selected_columns if col in ALL_COLUMNS])

        for lead in leads:
            row = []
            for key in selected_columns:
                if key not in ALL_COLUMNS:
                    continue
                    
                if key == "wa_link":
                    wa = lead.get("whatsapp", "") or lead.get("phone", "")
                    digits = "".join(filter(str.isdigit, wa))
                    if digits:
                        if not digits.startswith("55") and (10 <= len(digits) <= 11):
                            digits = "55" + digits
                        val = f"https://wa.me/{digits}"
                    else:
                        val = ""
                elif key in ("site_verified", "has_ads", "has_pixel", "has_analytics", "ssl_expired", "is_responsive"):
                    val = "Sim" if lead.get(key) else "Não"
                else:
                    val = lead.get(key, "")
                
                row.append(val)
            writer.writerow(row)


def export_xlsx(leads, filepath, selected_columns=None):
    if not selected_columns:
        selected_columns = list(ALL_COLUMNS.keys())

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(name="Inter", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    # Header
    actual_cols = [col for col in selected_columns if col in ALL_COLUMNS]
    for col_idx, key in enumerate(actual_cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=ALL_COLUMNS[key])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    data_font = Font(name="Inter", size=10)
    even_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    hot_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    warm_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

    for row_idx, lead in enumerate(leads, 2):
        for col_idx, key in enumerate(actual_cols, 1):
            if key == "wa_link":
                wa = lead.get("whatsapp", "") or lead.get("phone", "")
                digits = "".join(filter(str.isdigit, wa))
                if digits:
                    if not digits.startswith("55") and (10 <= len(digits) <= 11):
                        digits = "55" + digits
                    val = f"https://wa.me/{digits}"
                else:
                    val = ""
            elif key in ("site_verified", "has_ads", "has_pixel", "has_analytics", "ssl_expired", "is_responsive"):
                val = "Sim" if lead.get(key) else "Não"
            else:
                val = lead.get(key, "")

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

            if key == "tier":
                tier = lead.get("tier", "")
                if tier == "hot":
                    cell.fill = hot_fill
                elif tier == "warm":
                    cell.fill = warm_fill
            elif row_idx % 2 == 0:
                cell.fill = even_fill

    # Dynamic column widths
    for i in range(1, len(actual_cols) + 1):
        col_letter = chr(64 + i) if i <= 26 else f"A{chr(64 + i - 26)}"
        ws.column_dimensions[col_letter].width = 25

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filepath)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(filepath)
